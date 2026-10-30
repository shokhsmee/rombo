"""add-list-tools: openspec/changes/add-list-tools/specs/platform/spec.md at API level."""

import io

import pytest
from openpyxl import Workbook, load_workbook

from apps.accounting.seed import seed_chart
from apps.clients.models import Client
from apps.inventory.models import StockMove
from apps.procurement.tests import _project, _user_with
from conftest import API

INV = f"{API}/inventory"


@pytest.fixture(autouse=True)
def _chart(db):
    seed_chart()


def _xlsx(rows) -> io.BytesIO:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "import.xlsx"
    return buf


def _sheet(resp) -> list[tuple]:
    assert resp.status_code == 200, resp.content[:300]
    assert resp["Content-Type"].startswith("application/vnd.openxmlformats")
    return list(load_workbook(io.BytesIO(resp.content)).active.iter_rows(values_only=True))


def _client(admin, name, region):
    r = admin.post(f"{API}/clients/", {"name": name, "region": region}, format="json")
    assert r.status_code == 201, r.content
    return r.json()


def test_export_follows_the_filter(admin):
    _client(admin, "Ali", "Toshkent")
    _client(admin, "Vali", "Toshkent")
    _client(admin, "Sami", "Samarqand")
    rows = _sheet(admin.get(f"{API}/clients/export/", {"region": "Toshkent"}))
    assert rows[0][:3] == ("Kod", "Turi", "Nomi")
    assert sorted(r[2] for r in rows[1:]) == ["Ali", "Vali"]


def test_export_selected_ids(admin):
    a = _client(admin, "Ali", "Toshkent")
    _client(admin, "Vali", "Toshkent")
    rows = _sheet(admin.get(f"{API}/clients/export/", {"ids": [a["id"]]}))
    assert [r[2] for r in rows[1:]] == ["Ali"]


def test_template_has_required_marks(admin):
    rows = _sheet(admin.get(f"{API}/clients/import-template/"))
    assert "Nomi *" in rows[0] and "Telefon" in rows[0]


def _import(admin, rows, dry=False, url=f"{API}/clients/import/"):
    r = admin.post(url, {"file": _xlsx(rows), "dry_run": "true" if dry else "false"}, format="multipart")
    assert r.status_code == 200, r.content
    return r.json()


ROWS = [
    ["Nomi *", "Telefon", "Hudud"],
    ["Aziz Karimov", "+998901112233", "Toshkent"],
    [None, "+998901112244", "Buxoro"],
    ["Mega Qurilish MChJ", "", "Samarqand"],
]


def test_bad_rows_do_not_block_good_ones(admin):
    out = _import(admin, ROWS)
    assert (out["created"], out["total"]) == (2, 3)
    assert [r["ok"] for r in out["rows"]] == [True, False, True]
    assert out["rows"][1]["row"] == 3 and "Nomi" in out["rows"][1]["errors"][0]
    assert set(Client.objects.values_list("name", flat=True)) == {"Aziz Karimov", "Mega Qurilish MChJ"}


def test_dry_run_saves_nothing(admin):
    out = _import(admin, ROWS, dry=True)
    assert (out["created"], out["dry_run"]) == (2, True)
    assert not Client.objects.exists()


def test_import_resolves_references(admin):
    _client(admin, "Ali Valiyev", "Toshkent")
    out = _import(
        admin,
        [
            ["Nomi", "Mijoz", "Summa", "Sotuv sanasi"],
            ["Oshxona", "ali valiyev", 12500000, "31.10.2026"],
            ["Ofis", "Yo'q mijoz", 1, None],
        ],
        url=f"{API}/projects/import/",
    )
    assert [r["ok"] for r in out["rows"]] == [True, False]
    assert "topilmadi" in out["rows"][1]["errors"][0]


def test_import_needs_create_permission(api, admin):
    viewer = _user_with(api, admin, "kuzatuvchi", {"clients": {"can_view": True}})
    r = viewer.post(f"{API}/clients/import/", {"file": _xlsx(ROWS)}, format="multipart")
    assert r.status_code == 403
    assert viewer.get(f"{API}/clients/tabular/").json()["import"] is None


def _stocked_drafts(admin, qtys):
    p = admin.post(
        f"{INV}/products/", {"name": "Profil", "unit": "m", "category": "profil"}, format="json"
    ).json()
    m = admin.post(
        f"{INV}/moves/",
        {"type": "adjust", "product_id": p["id"], "qty": "10", "unit_cost": "1000"},
        format="json",
    ).json()
    admin.post(f"{INV}/moves/{m['id']}/post/")
    project = _project(admin)
    return [
        admin.post(
            f"{INV}/moves/",
            {"type": "issue", "product_id": p["id"], "qty": q, "project_id": project["id"]},
            format="json",
        ).json()["id"]
        for q in qtys
    ]


def test_partial_success(admin):
    ids = _stocked_drafts(admin, ["4", "4", "4"])
    r = admin.post(f"{INV}/moves/bulk/", {"action": "post", "ids": ids}, format="json")
    assert r.status_code == 200, r.content
    out = r.json()
    assert out["done"] == 2 and len(out["failed"]) == 1
    assert "Omborda yetarli emas" in out["failed"][0]["message"]
    assert StockMove.objects.filter(pk__in=ids, status="posted").count() == 2


def test_bulk_permissions_and_unknown_action(api, admin):
    ids = _stocked_drafts(admin, ["1"])
    viewer = _user_with(api, admin, "kuzatuvchi2", {"inventory": {"can_view": True}})
    assert viewer.post(f"{INV}/moves/bulk/", {"action": "post", "ids": ids}, format="json").status_code == 403
    assert viewer.get(f"{INV}/moves/tabular/").json()["bulk"] == []
    assert admin.post(f"{INV}/moves/bulk/", {"action": "nope", "ids": ids}, format="json").status_code == 400
    keys = [b["key"] for b in admin.get(f"{INV}/moves/tabular/").json()["bulk"]]
    assert keys == ["post", "delete"]


def test_bulk_with_parameter_and_comment(admin):
    a, b = _project(admin, status="draft"), _project(admin, status="draft")
    r = admin.post(
        f"{API}/projects/bulk/",
        {"action": "status", "ids": [a["id"], b["id"]], "params": {"status": "confirmed"}},
        format="json",
    )
    assert r.json()["done"] == 2
    assert admin.get(f"{API}/projects/{a['id']}/").json()["status"] == "confirmed"
    r = admin.post(f"{API}/procurement/requests/bulk/", {"action": "cancel", "ids": [1]}, format="json")
    assert r.status_code == 400 and r.json()["error"]["field"] == "comment"


@pytest.mark.parametrize(
    "url",
    [
        "/projects/",
        "/clients/",
        "/employees/",
        "/users/",
        "/access/roles/",
        "/pricing/",
        "/procurement/requests/",
        "/procurement/orders/",
        "/procurement/receipts/",
        "/procurement/suppliers/",
        "/accounting/invoices/",
        "/accounting/entries/",
        "/kassa/orders/",
        "/kassa/payment-requests/",
        "/inventory/moves/",
        "/inventory/products/",
        "/telephony/calls/",
    ],
)
def test_every_list_exports(admin, url):
    meta = admin.get(f"{API}{url}tabular/")
    assert meta.status_code == 200, (url, meta.content[:200])
    assert len(_sheet(admin.get(f"{API}{url}export/"))) >= 1


def test_stock_exports(admin):
    _stocked_drafts(admin, [])
    rows = _sheet(admin.get(f"{INV}/stock/", {"xlsx": "1"}))
    assert rows[0][0] == "Kod" and rows[1][1] == "Profil"
