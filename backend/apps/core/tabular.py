"""List tools (add-list-tools, ADR 0011): Excel export, Excel import, batch operations — once, for every list.

A viewset opts in by mixing `TabularMixin` in front of `ServiceModelViewSet` and declaring:

    excel_name = "mijozlar"                      # file name stem
    excel_columns = [Col("Kod", "code"), Col("Nomi", "name"), Col("Summa", "amount", kind="money"), ...]
    excel_import = [Field("Nomi", "name", required=True), Field("Mijoz", "client_id", resolve=ref(Client)), ...]
    bulk_actions = [Bulk("delete", "O'chirish", services.thing_delete, action=Action.DELETE, danger=True), ...]

Endpoints (detail=False): `tabular/` (what the page may offer), `export/`, `import-template/`, `import/`,
`bulk/`. An imported row goes through the viewset's serializer and `service_create`; a batch action calls
the same service as the single-record button. Every row / id runs in its own savepoint.
"""

from __future__ import annotations

import datetime as dt
import io
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from apps.access.registry import Action
from apps.access.services import has_permission
from apps.core.exceptions import DomainError, ValidationFailed

MAX_EXPORT = 10_000
MAX_IMPORT = 2_000
MAX_BULK = 500
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ------------------------------------------------------------------ declarations


@dataclass(frozen=True)
class Col:
    """One export column: `get` is an attribute path ("client.name") or a function of the row."""

    header: str
    get: str | Callable[[Any], Any]
    kind: str = ""  # money | qty | date — number formats
    width: int = 18

    def value(self, obj):
        if callable(self.get):
            return self.get(obj)
        v = obj
        for part in self.get.split("."):
            v = getattr(v, part, None) if v is not None else None
        return v() if callable(v) else v


@dataclass(frozen=True)
class Field:
    """One import column. `resolve` turns the cell into the serializer value (raise ValueError to reject)."""

    header: str
    name: str  # serializer input key
    required: bool = False
    resolve: Callable[[Any], Any] | None = None
    hint: str = ""

    def parse(self, raw):
        if self.resolve is not None:
            return self.resolve(raw)
        if isinstance(raw, dt.datetime):
            return raw.date().isoformat()
        if isinstance(raw, dt.date):
            return raw.isoformat()
        if isinstance(raw, float) and raw.is_integer():
            return str(int(raw))
        return str(raw).strip() if not isinstance(raw, bool) else raw


@dataclass(frozen=True)
class Bulk:
    """A batch action: `fn(obj, *, actor, **params)` — the same service as the single-record button."""

    key: str
    label: str
    fn: Callable
    action: Action = Action.UPDATE
    page: str | None = None  # another page's permission (e.g. payment_approve)
    danger: bool = False
    comment: bool = False  # asks for a reason, passed as `comment`
    params: list[dict] = field(default_factory=list)  # [{"name", "label", "options": [{"value","label"}]}]
    only: Callable[[Any], bool] | None = None  # rows it applies to (the page greys out the rest)


# ------------------------------------------------------------------ cell parsers (for Field.resolve)


def as_decimal(raw) -> str:
    if isinstance(raw, int | float | Decimal):
        return str(Decimal(str(raw)))
    s = str(raw).replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return str(Decimal(s))
    except InvalidOperation as e:
        raise ValueError(f"«{raw}» son emas") from e


def as_date(raw) -> str:
    if isinstance(raw, dt.datetime):
        return raw.date().isoformat()
    if isinstance(raw, dt.date):
        return raw.isoformat()
    s = str(raw).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"«{raw}» sana emas (31.10.2026)")


def as_bool(raw) -> bool:
    if isinstance(raw, bool):
        return raw
    s = str(raw).strip().lower()
    if s in ("1", "ha", "yes", "true", "+", "faol", "bor"):
        return True
    if s in ("0", "yo'q", "yoq", "no", "false", "-", "nofaol", ""):
        return False
    raise ValueError(f"«{raw}»: ha / yo'q")


def choice(choices) -> Callable[[Any], str]:
    """Accepts the stored value or its Uzbek label ("Tayyor mahsulot" → "tayyor")."""

    def parse(raw):
        s = str(raw).strip().lower()
        for value, label in choices:
            if s in (str(value).lower(), str(label).lower()):
                return value
        raise ValueError(f"«{raw}»: " + " / ".join(str(label) for _, label in choices))

    return parse


def ref(model, *fields: str, extra: Q | None = None) -> Callable[[Any], int]:
    """A reference by code or name (or any of `fields`), case-insensitive → its pk."""
    fields = fields or ("code", "name")

    def parse(raw):
        s = str(raw).strip()
        if isinstance(raw, float) and raw.is_integer():
            s = str(int(raw))
        qs = model._default_manager.all()
        if extra is not None:
            qs = qs.filter(extra)
        for f in fields:
            hits = list(qs.filter(**{f"{f}__iexact": s}).values_list("pk", flat=True)[:2])
            if len(hits) == 1:
                return hits[0]
            if len(hits) > 1:
                raise ValueError(f"«{s}» bir nechta — kodini yozing")
        raise ValueError(f"«{s}» topilmadi")

    return parse


# ------------------------------------------------------------------ xlsx


def _cell(v):
    if isinstance(v, dt.datetime):
        return timezone.localtime(v).replace(tzinfo=None) if timezone.is_aware(v) else v
    if isinstance(v, bool):
        return "Ha" if v else "Yo'q"
    if isinstance(v, Decimal | int | float | dt.date | str) or v is None:
        return v
    return str(v)


def workbook(
    title: str,
    headers: list[str],
    rows,
    kinds: list[str] | None = None,
    widths: list[int] | None = None,
    notes: list[str] | None = None,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E2ECF8")
        ws.column_dimensions[get_column_letter(i)].width = (widths or [18] * len(headers))[i - 1]
        if notes and notes[i - 1]:
            c.comment = Comment(notes[i - 1], "Rombo")
    formats = {"money": "#,##0.00", "qty": "#,##0.###", "date": "DD.MM.YYYY"}
    for row in rows:
        ws.append([_cell(v) for v in row])
    if kinds:
        for i, k in enumerate(kinds, start=1):
            if k in formats:
                for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell.number_format = formats[k]
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(data: bytes, stem: str) -> HttpResponse:
    resp = HttpResponse(data, content_type=XLSX)
    resp["Content-Disposition"] = f'attachment; filename="{stem}-{timezone.localdate():%Y-%m-%d}.xlsx"'
    return resp


def _norm(h) -> str:
    return str(h or "").replace("*", "").strip().lower()


def _flatten(errors, labels: dict[str, str]) -> list[str]:
    out = []
    if isinstance(errors, dict):
        for k, v in errors.items():
            msgs = v if isinstance(v, list) else [v]
            for m in msgs:
                text = "; ".join(_flatten(m, labels)) if isinstance(m, dict | list) else str(m)
                out.append(f"{labels.get(k, k)}: {text}" if k != "non_field_errors" else text)
    elif isinstance(errors, list):
        for m in errors:
            out.extend(_flatten(m, labels))
    else:
        out.append(str(errors))
    return out


# ------------------------------------------------------------------ the mixin


class TabularMixin:
    excel_name = "royxat"
    excel_columns: list[Col] = []
    excel_import: list[Field] | None = None
    import_defaults: dict = {}
    bulk_actions: list[Bulk] = []

    def _can(self, action_: Action, page: str | None = None) -> bool:
        return has_permission(self.request.user, page or self.page_key, action_)

    @action(detail=False, methods=["get"])
    def tabular(self, request):
        """What the list page may offer this user."""
        can_import = self.excel_import is not None and self._can(Action.CREATE)
        return Response(
            {
                "export": bool(self.excel_columns),
                "import": [
                    {"header": f.header, "required": f.required, "hint": f.hint} for f in self.excel_import
                ]
                if can_import
                else None,
                "bulk": [
                    {
                        "key": b.key,
                        "label": b.label,
                        "danger": b.danger,
                        "comment": b.comment,
                        "params": b.params,
                    }
                    for b in self.bulk_actions
                    if self._can(b.action, b.page)
                ],
            }
        )

    @action(detail=False, methods=["get"])
    def export(self, request):
        if not self.excel_columns:
            raise ValidationFailed("Bu ro'yxat Excelga chiqarilmaydi")
        qs = self.filter_queryset(self.get_queryset())
        ids = [int(i) for i in request.query_params.getlist("ids") if str(i).isdigit()]
        if ids:
            qs = qs.filter(pk__in=ids)
        cols = self.excel_columns
        rows = ([c.value(obj) for c in cols] for obj in qs[:MAX_EXPORT])
        data = workbook(
            self.excel_name, [c.header for c in cols], rows, [c.kind for c in cols], [c.width for c in cols]
        )
        return xlsx_response(data, self.excel_name)

    @action(detail=False, methods=["get"], url_path="import-template")
    def import_template(self, request):
        if self.excel_import is None:
            raise ValidationFailed("Bu ro'yxatga Excel import yo'q")
        fields_ = self.excel_import
        data = workbook(
            self.excel_name,
            [f"{f.header}{' *' if f.required else ''}" for f in fields_],
            [],
            widths=[max(14, len(f.header) + 4) for f in fields_],
            notes=[f.hint for f in fields_],
        )
        return xlsx_response(data, f"{self.excel_name}-shablon")

    @action(detail=False, methods=["post"], url_path="import", parser_classes=[MultiPartParser, FormParser])
    def import_rows(self, request):
        if self.excel_import is None:
            raise ValidationFailed("Bu ro'yxatga Excel import yo'q")
        if not self._can(Action.CREATE):
            raise PermissionDenied("Bu amal uchun ruxsat yo'q")
        upload = request.FILES.get("file")
        if upload is None:
            raise ValidationFailed("Excel faylni tanlang (.xlsx)", field="file")
        dry = str(request.data.get("dry_run", "")).lower() in ("1", "true", "yes")
        try:
            from openpyxl import load_workbook

            ws = load_workbook(upload, read_only=True, data_only=True).worksheets[0]
            table = list(ws.iter_rows(values_only=True))
        except Exception as e:  # not an xlsx, broken file
            raise ValidationFailed(
                "Faylni o'qib bo'lmadi — .xlsx shablondan foydalaning", field="file"
            ) from e
        if not table:
            raise ValidationFailed("Fayl bo'sh", field="file")
        by_header = {_norm(f.header): f for f in self.excel_import}
        colmap = [(i, by_header[_norm(h)]) for i, h in enumerate(table[0]) if _norm(h) in by_header]
        present = {f.name for _, f in colmap}
        missing = [f.header for f in self.excel_import if f.required and f.name not in present]
        if missing:
            raise ValidationFailed("Ustun yo'q: " + ", ".join(missing), field="file")
        body = [(n, r) for n, r in enumerate(table[1:], start=2) if any(v not in (None, "") for v in r)]
        if len(body) > MAX_IMPORT:
            raise ValidationFailed(f"Bir faylda ko'pi bilan {MAX_IMPORT} qator", field="file")
        labels = {f.name: f.header for f in self.excel_import}
        results, created = [], 0
        with transaction.atomic():
            for n, row in body:
                payload, errors = dict(self.import_defaults), []
                for i, f in colmap:
                    raw = row[i] if i < len(row) else None
                    if raw is None or (isinstance(raw, str) and not raw.strip()):
                        if f.required:
                            errors.append(f"{f.header}: to'ldirilmagan")
                        continue
                    try:
                        payload[f.name] = f.parse(raw)
                    except ValueError as e:
                        errors.append(f"{f.header}: {e}")
                label = ""
                if not errors:
                    ser = self.get_serializer(data=payload)
                    if not ser.is_valid():
                        errors = _flatten(ser.errors, labels)
                    else:
                        try:
                            with transaction.atomic():
                                obj = type(self).service_create(actor=request.user, **ser.validated_data)
                            label, created = str(obj), created + 1
                        except DomainError as e:
                            errors.append(f"{labels.get(e.field, '')}: {e.message}".lstrip(": "))
                        except DjangoValidationError as e:
                            errors.extend(e.messages)
                results.append({"row": n, "ok": not errors, "errors": errors, "label": label})
            if dry:
                transaction.set_rollback(True)
        return Response({"rows": results, "created": created, "total": len(results), "dry_run": dry})

    @action(detail=False, methods=["post"])
    def bulk(self, request):
        key = request.data.get("action")
        act = next((b for b in self.bulk_actions if b.key == key), None)
        if act is None:
            raise ValidationFailed("Noma'lum amal", field="action")
        if not self._can(act.action, act.page):
            raise PermissionDenied("Bu amal uchun ruxsat yo'q")
        raw_ids = request.data.get("ids") or []
        if not isinstance(raw_ids, list) or not raw_ids:
            raise ValidationFailed("Qator tanlanmagan", field="ids")
        if len(raw_ids) > MAX_BULK:
            raise ValidationFailed(f"Bir martada ko'pi bilan {MAX_BULK} ta", field="ids")
        ids = [int(i) for i in raw_ids if str(i).isdigit()]
        params = dict(request.data.get("params") or {})
        allowed = {p["name"] for p in act.params} | ({"comment"} if act.comment else set())
        params = {k: v for k, v in params.items() if k in allowed}
        if act.comment and not str(params.get("comment", "")).strip():
            raise ValidationFailed("Sababini yozing", field="comment")
        qs = self.get_queryset()
        done, failed = [], []
        for pk in ids:
            obj = qs.filter(pk=pk).first()
            if obj is None:
                failed.append({"id": pk, "label": f"#{pk}", "message": "Topilmadi"})
                continue
            if act.only is not None and not act.only(obj):
                failed.append({"id": pk, "label": str(obj), "message": "Bu yozuvga qo'llanmaydi"})
                continue
            try:
                with transaction.atomic():
                    act.fn(obj, actor=request.user, **params)
                done.append(pk)
            except DomainError as e:
                failed.append({"id": pk, "label": str(obj), "message": e.message})
            except DjangoValidationError as e:
                failed.append({"id": pk, "label": str(obj), "message": "; ".join(e.messages)})
        return Response({"action": act.key, "done": len(done), "done_ids": done, "failed": failed})


def opts(choices) -> list[dict]:
    """Options of a batch-action parameter from a TextChoices class."""
    return [{"value": v, "label": label} for v, label in choices.choices]


def label_of(choices) -> Callable[[Any], str]:
    """Export helper: the Uzbek label of a stored choice ("tayyor" → "Tayyor mahsulot")."""
    names = dict(choices.choices)
    return lambda v: names.get(v, v)
